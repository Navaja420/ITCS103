import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

file_name = "Navaja_Database.xlsx"

# Create Excel file if not exists
if not os.path.exists(file_name):
    wb = Workbook()
    ws = wb.active
    ws.append(["ID", "Title", "Author", "Year", "Status"])
    wb.save(file_name)

# Load data
def load_data():
    for row in tree.get_children():
        tree.delete(row)

    wb = load_workbook(file_name)
    ws = wb.active

    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        tree.insert("", tk.END, values=row)
        count += 1

    lbl_total.config(text=f"Total Books: {count}")

# Generate ID
def generate_id():
    wb = load_workbook(file_name)
    ws = wb.active
    return ws.max_row

# Add Book
def add_book():
    title = entry_title.get()
    author = entry_author.get()
    year = entry_year.get()
    status = status_var.get()

    if title == "" or author == "" or year == "":
        messagebox.showerror("Error", "All fields are required!")
        return

    if not year.isdigit():
        messagebox.showerror("Error", "Year must be a number!")
        return

    wb = load_workbook(file_name)
    ws = wb.active

    new_id = generate_id()
    ws.append([new_id, title, author, year, status])
    wb.save(file_name)

    messagebox.showinfo("Success", "Book Added Successfully!")
    clear_fields()
    load_data()

# Select Book
def select_book(event):
    selected = tree.focus()
    data = tree.item(selected, "values")

    if data:
        entry_id.delete(0, tk.END)
        entry_id.insert(0, data[0])

        entry_title.delete(0, tk.END)
        entry_title.insert(0, data[1])

        entry_author.delete(0, tk.END)
        entry_author.insert(0, data[2])

        entry_year.delete(0, tk.END)
        entry_year.insert(0, data[3])

        status_var.set(data[4])

# Update Book
def update_book():
    selected_id = entry_id.get()

    if selected_id == "":
        messagebox.showerror("Error", "Select a record first!")
        return

    wb = load_workbook(file_name)
    ws = wb.active

    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == selected_id:
            row[1].value = entry_title.get()
            row[2].value = entry_author.get()
            row[3].value = entry_year.get()
            row[4].value = status_var.get()

    wb.save(file_name)
    messagebox.showinfo("Updated", "Book Updated Successfully!")
    load_data()

# Delete Book
def delete_book():
    selected_id = entry_id.get()

    if selected_id == "":
        messagebox.showerror("Error", "Select a record first!")
        return

    confirm = messagebox.askyesno("Confirm", "Are you sure you want to delete?")
    if not confirm:
        return

    wb = load_workbook(file_name)
    ws = wb.active

    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == selected_id:
            ws.delete_rows(row[0].row)

    wb.save(file_name)
    messagebox.showinfo("Deleted", "Book Deleted Successfully!")
    clear_fields()
    load_data()

# Clear fields
def clear_fields():
    entry_id.delete(0, tk.END)
    entry_title.delete(0, tk.END)
    entry_author.delete(0, tk.END)
    entry_year.delete(0, tk.END)
    status_var.set("Available")

# GUI
root = tk.Tk()
root.title("Library Management System")
root.geometry("800x500")

# Labels and Entries
tk.Label(root, text="ID").grid(row=0, column=0, padx=5, pady=5)
entry_id = tk.Entry(root)
entry_id.grid(row=0, column=1)

tk.Label(root, text="Title").grid(row=1, column=0)
entry_title = tk.Entry(root)
entry_title.grid(row=1, column=1)

tk.Label(root, text="Author").grid(row=2, column=0)
entry_author = tk.Entry(root)
entry_author.grid(row=2, column=1)

tk.Label(root, text="Year").grid(row=3, column=0)
entry_year = tk.Entry(root)
entry_year.grid(row=3, column=1)

tk.Label(root, text="Status").grid(row=4, column=0)
status_var = tk.StringVar(value="Available")
status_menu = ttk.Combobox(root, textvariable=status_var, values=["Available", "Borrowed"])
status_menu.grid(row=4, column=1)

# Buttons
tk.Button(root, text="Add Book", width=12, command=add_book).grid(row=5, column=0, pady=10)
tk.Button(root, text="Update", width=12, command=update_book).grid(row=5, column=1)
tk.Button(root, text="Delete", width=12, command=delete_book).grid(row=5, column=2)
tk.Button(root, text="Clear", width=12, command=clear_fields).grid(row=5, column=3)

# Table
tree = ttk.Treeview(root, columns=("ID", "Title", "Author", "Year", "Status"), show="headings")

for col in ("ID", "Title", "Author", "Year", "Status"):
    tree.heading(col, text=col)
    tree.column(col, width=120)

tree.grid(row=6, column=0, columnspan=4, pady=20)
tree.bind("<ButtonRelease-1>", select_book)

# Total Books Label
lbl_total = tk.Label(root, text="Total Books: 0", font=("Arial", 12, "bold"))
lbl_total.grid(row=7, column=0, columnspan=2)

# Load Data
load_data()

root.mainloop()