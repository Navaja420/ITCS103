import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

# Library Management System
# Created by: [Your Name]
# Course: BSIT 1st Year

class LibraryManagementSystem:
    def __init__(self, root):
        self.root = root
        self.root.title("Library Management System")
        self.root.geometry("1100x650")
        self.root.configure(bg="brown")
        
        # Excel file name
        self.excel_file = "Navaja_Database.xlsx"
        
        # Create excel file if not exist
        self.create_excel_file()
        
        # Variables for input
        self.book_id = tk.StringVar()
        self.book_title = tk.StringVar()
        self.author = tk.StringVar()
        self.isbn = tk.StringVar()
        self.publisher = tk.StringVar()
        self.year = tk.StringVar()
        self.category = tk.StringVar()
        self.quantity = tk.StringVar()
        self.available = tk.StringVar()
        
        # Call function to create widgets
        self.create_widgets()
        self.load_data()
    
    def create_excel_file(self):
        """Create excel file"""
        if not os.path.exists(self.excel_file):
            wb = Workbook()
            ws = wb.active
            ws.title = "Books"
            # Headers
            headers = ["ID", "Book Title", "Author", "ISBN", "Publisher", 
                      "Year", "Category", "Quantity", "Available"]
            ws.append(headers)
            wb.save(self.excel_file)
    
    def create_widgets(self):
        """Create GUI"""
        # Title
        title = tk.Label(self.root, text="LIBRARY MANAGEMENT SYSTEM", 
                        font=("Arial", 20, "bold"), bg="brown")
        title.pack(pady=10)
        
        # Frame for inputs
        frame = tk.Frame(self.root, bg="brown", bd=2, relief=tk.GROOVE)
        frame.pack(pady=10, padx=10, fill=tk.X)
        
        # Row 1 - Book Title, Author, ISBN
        tk.Label(frame, text="Book Title:", font=("Arial", 10), 
                bg="lightblue").grid(row=0, column=0, padx=5, pady=5, sticky="e")
        tk.Entry(frame, textvariable=self.book_title, width=25).grid(row=0, column=1, padx=5, pady=5)
        
        tk.Label(frame, text="Author:", font=("Arial", 10), 
                bg="lightblue").grid(row=0, column=2, padx=5, pady=5, sticky="e")
        tk.Entry(frame, textvariable=self.author, width=25).grid(row=0, column=3, padx=5, pady=5)
        
        tk.Label(frame, text="ISBN:", font=("Arial", 10), 
                bg="lightblue").grid(row=0, column=4, padx=5, pady=5, sticky="e")
        tk.Entry(frame, textvariable=self.isbn, width=20).grid(row=0, column=5, padx=5, pady=5)
        
        # Row 2 - Publisher, Year, Category
        tk.Label(frame, text="Publisher:", font=("Arial", 10), 
                bg="lightblue").grid(row=1, column=0, padx=5, pady=5, sticky="e")
        tk.Entry(frame, textvariable=self.publisher, width=25).grid(row=1, column=1, padx=5, pady=5)
        
        tk.Label(frame, text="Year:", font=("Arial", 10), 
                bg="lightblue").grid(row=1, column=2, padx=5, pady=5, sticky="e")
        tk.Entry(frame, textvariable=self.year, width=25).grid(row=1, column=3, padx=5, pady=5)
        
        tk.Label(frame, text="Category:", font=("Arial", 10), 
                bg="lightblue").grid(row=1, column=4, padx=5, pady=5, sticky="e")
        self.category_combo = ttk.Combobox(frame, textvariable=self.category, width=18, state="readonly")
        self.category_combo['values'] = ("Fiction", "Non-Fiction", "Science", "Technology", "History", "Other")
        self.category_combo.grid(row=1, column=5, padx=5, pady=5)
        
        # Row 3 - Quantity, Available
        tk.Label(frame, text="Quantity:", font=("Arial", 10), 
                bg="lightblue").grid(row=2, column=0, padx=5, pady=5, sticky="e")
        tk.Entry(frame, textvariable=self.quantity, width=25).grid(row=2, column=1, padx=5, pady=5)
        
        tk.Label(frame, text="Available:", font=("Arial", 10), 
                bg="lightblue").grid(row=2, column=2, padx=5, pady=5, sticky="e")
        tk.Entry(frame, textvariable=self.available, width=25).grid(row=2, column=3, padx=5, pady=5)
        
        # Buttons
        btn_frame = tk.Frame(self.root, bg="lightblue")
        btn_frame.pack(pady=20)
        
        tk.Button(btn_frame, text="ADD BOOK", command=self.add_book, 
                 bg="green", fg="white", font=("Arial", 10, "bold"), 
                 width=12).grid(row=0, column=0, padx=10)
        
        tk.Button(btn_frame, text="UPDATE", command=self.update_book, 
                 bg="blue", fg="white", font=("Arial", 10, "bold"), 
                 width=12).grid(row=0, column=1, padx=10)
        
        tk.Button(btn_frame, text="DELETE", command=self.delete_book, 
                 bg="red", fg="white", font=("Arial", 10, "bold"), 
                 width=12).grid(row=0, column=2, padx=10)
        
        tk.Button(btn_frame, text="CLEAR", command=self.clear_form, 
                 bg="gray", fg="white", font=("Arial", 10, "bold"), 
                 width=12).grid(row=0, column=3, padx=10)
        
        tk.Button(btn_frame, text="REFRESH", command=self.load_data, 
                 bg="teal", fg="white", font=("Arial", 10, "bold"), 
                 width=12).grid(row=0, column=4, padx=10)
        
        # Table
        tree_frame = tk.Frame(self.root)
        tree_frame.pack(pady=10, padx=10, fill=tk.BOTH, expand=True)
        
        # Scrollbar
        y_scroll = ttk.Scrollbar(tree_frame, orient=tk.VERTICAL)
        
        # Treeview
        columns = ("ID", "Book Title", "Author", "ISBN", "Publisher", 
                  "Year", "Category", "Quantity", "Available")
        
        self.tree = ttk.Treeview(tree_frame, columns=columns, 
                                yscrollcommand=y_scroll.set)
        
        y_scroll.config(command=self.tree.yview)
        
        # Headings
        self.tree.heading("ID", text="ID")
        self.tree.heading("Book Title", text="Book Title")
        self.tree.heading("Author", text="Author")
        self.tree.heading("ISBN", text="ISBN")
        self.tree.heading("Publisher", text="Publisher")
        self.tree.heading("Year", text="Year")
        self.tree.heading("Category", text="Category")
        self.tree.heading("Quantity", text="Quantity")
        self.tree.heading("Available", text="Available")
        
        # Column widths
        self.tree.column("ID", width=50)
        self.tree.column("Book Title", width=150)
        self.tree.column("Author", width=120)
        self.tree.column("ISBN", width=100)
        self.tree.column("Publisher", width=120)
        self.tree.column("Year", width=60)
        self.tree.column("Category", width=100)
        self.tree.column("Quantity", width=70)
        self.tree.column("Available", width=80)
        
        self.tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        y_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        
        # Bind click event
        self.tree.bind("<<TreeviewSelect>>", self.on_select)
    
    def get_next_id(self):
        """Get next ID number"""
        wb = load_workbook(self.excel_file)
        ws = wb.active
        max_id = 0
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            if row[0] is not None:
                if row[0] > max_id:
                    max_id = row[0]
        return max_id + 1
    
    def check_inputs(self):
        """Check if inputs are not empty"""
        if self.book_title.get() == "":
            messagebox.showerror("Error", "Book Title is required!")
            return False
        if self.author.get() == "":
            messagebox.showerror("Error", "Author is required!")
            return False
        if self.isbn.get() == "":
            messagebox.showerror("Error", "ISBN is required!")
            return False
        if self.publisher.get() == "":
            messagebox.showerror("Error", "Publisher is required!")
            return False
        if self.year.get() == "":
            messagebox.showerror("Error", "Year is required!")
            return False
        if self.category.get() == "":
            messagebox.showerror("Error", "Category is required!")
            return False
        if self.quantity.get() == "":
            messagebox.showerror("Error", "Quantity is required!")
            return False
        if self.available.get() == "":
            messagebox.showerror("Error", "Available copies is required!")
            return False
        
        # Check if numbers are valid
        try:
            int(self.year.get())
            int(self.quantity.get())
            int(self.available.get())
        except:
            messagebox.showerror("Error", "Year, Quantity, and Available must be numbers!")
            return False
        
        return True
    
    def add_book(self):
        """Add new book"""
        # Check inputs first
        if not self.check_inputs():
            return
        
        # Get next ID
        new_id = self.get_next_id()
        
        # Prepare data
        data = [
            new_id,
            self.book_title.get(),
            self.author.get(),
            self.isbn.get(),
            self.publisher.get(),
            self.year.get(),
            self.category.get(),
            self.quantity.get(),
            self.available.get()
        ]
        
        # Save to excel
        try:
            wb = load_workbook(self.excel_file)
            ws = wb.active
            ws.append(data)
            wb.save(self.excel_file)
            messagebox.showinfo("Success", "Book added successfully! ID: " + str(new_id))
            self.clear_form()
            self.load_data()
        except Exception as e:
            messagebox.showerror("Error", "Failed to add book: " + str(e))
    
    def load_data(self):
        """Load all books from excel"""
        # Clear table
        for item in self.tree.get_children():
            self.tree.delete(item)
        
        # Load from excel
        try:
            wb = load_workbook(self.excel_file)
            ws = wb.active
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is not None:
                    self.tree.insert("", tk.END, values=row)
        except Exception as e:
            messagebox.showerror("Error", "Failed to load data: " + str(e))
    
    def on_select(self, event):
        """When user clicks on table"""
        selected = self.tree.selection()
        if not selected:
            return
        
        item = self.tree.item(selected[0])
        values = item['values']
        
        # Fill the inputs
        self.book_id.set(values[0])
        self.book_title.set(values[1])
        self.author.set(values[2])
        self.isbn.set(values[3])
        self.publisher.set(values[4])
        self.year.set(values[5])
        self.category.set(values[6])
        self.quantity.set(values[7])
        self.available.set(values[8])
    
    def update_book(self):
        """Update existing book"""
        # Check if selected
        selected = self.tree.selection()
        if not selected:
            messagebox.showerror("Error", "Please select a book to update!")
            return
        
        # Check inputs
        if not self.check_inputs():
            return
        
        book_id = self.book_id.get()
        if book_id == "":
            messagebox.showerror("Error", "No book selected!")
            return
        
        # Ask confirmation
        confirm = messagebox.askyesno("Confirm", "Update this book?")
        if not confirm:
            return
        
        # Prepare data
        data = [
            int(book_id),
            self.book_title.get(),
            self.author.get(),
            self.isbn.get(),
            self.publisher.get(),
            self.year.get(),
            self.category.get(),
            self.quantity.get(),
            self.available.get()
        ]
        
        # Update in excel
        try:
            wb = load_workbook(self.excel_file)
            ws = wb.active
            
            # Find the row
            for row in ws.iter_rows(min_row=2):
                if row[0].value == int(book_id):
                    for i, cell in enumerate(row):
                        cell.value = data[i]
                    break
            
            wb.save(self.excel_file)
            messagebox.showinfo("Success", "Book updated successfully!")
            self.clear_form()
            self.load_data()
        except Exception as e:
            messagebox.showerror("Error", "Failed to update: " + str(e))
    
    def delete_book(self):
        """Delete a book"""
        # Check if selected
        selected = self.tree.selection()
        if not selected:
            messagebox.showerror("Error", "Please select a book to delete!")
            return
        
        book_id = self.book_id.get()
        if book_id == "":
            messagebox.showerror("Error", "No book selected!")
            return
        
        # Ask confirmation
        confirm = messagebox.askyesno("Confirm", "Delete Book ID: " + str(book_id) + "?")
        if not confirm:
            return
        
        # Delete from excel
        try:
            wb = load_workbook(self.excel_file)
            ws = wb.active
            
            # Find and delete row
            for row in ws.iter_rows(min_row=2):
                if row[0].value == int(book_id):
                    ws.delete_rows(row[0].row, 1)
                    break
            
            wb.save(self.excel_file)
            messagebox.showinfo("Success", "Book deleted successfully!")
            self.clear_form()
            self.load_data()
        except Exception as e:
            messagebox.showerror("Error", "Failed to delete: " + str(e))
    
    def clear_form(self):
        """Clear all inputs"""
        self.book_id.set("")
        self.book_title.set("")
        self.author.set("")
        self.isbn.set("")
        self.publisher.set("")
        self.year.set("")
        self.category.set("")
        self.quantity.set("")
        self.available.set("")

# Main program
if __name__ == "__main__":
    root = tk.Tk()
    app = LibraryManagementSystem(root)
    root.mainloop()
